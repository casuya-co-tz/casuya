from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session

from backend.config.database import get_db, redis_client
from backend.middleware.auth import get_current_user
from backend.middleware.permissions import require_role
from backend.models.student import Student
from backend.models.teacher import Teacher
from backend.models.user import User
from backend.schemas.users import UserResponse, UserStatusUpdateRequest, UserUpdateRequest

router = APIRouter(prefix="/users", tags=["users"])

# MIME type for .xlsx files
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

EXPORT_HEADERS = [
    "Name",
    "Email",
    "Phone",
    "Role",
    "Status",
    "Form Level",
    "Subjects",
    "Registered (UTC)",
]


def _serialize_user(u: User) -> dict:
    return {
        "id": u.id,
        "email": u.email,
        "full_name": u.full_name,
        "phone": u.phone,
        "role": u.role,
        "is_active": u.is_active,
        "created_at": u.created_at.isoformat() if u.created_at else None,
    }


def _attach_profiles(users: list[dict], db: Session) -> None:
    """Attach student/teacher profile metadata to each serialized user dict."""
    if not users:
        return
    user_ids = [u["id"] for u in users]
    students = (
        db.query(Student.user_id, Student.id, Student.full_name, Student.form_level)
        .filter(Student.user_id.in_(user_ids))
        .all()
    )
    teachers = (
        db.query(Teacher.user_id, Teacher.id, Teacher.full_name, Teacher.subjects)
        .filter(Teacher.user_id.in_(user_ids))
        .all()
    )
    student_map = {row.user_id: row for row in students}
    teacher_map = {row.user_id: row for row in teachers}
    for item in users:
        uid = item["id"]
        student = student_map.get(uid)
        teacher = teacher_map.get(uid)
        if student is not None:
            item["profile"] = {
                "type": "student",
                "id": student.id,
                "form_level": student.form_level,
            }
            if not item["full_name"]:
                item["full_name"] = student.full_name
        elif teacher is not None:
            item["profile"] = {
                "type": "teacher",
                "id": teacher.id,
                "subjects": teacher.subjects,
            }
            if not item["full_name"]:
                item["full_name"] = teacher.full_name


def _invalidate_user_cache(user_id: str) -> None:
    try:
        redis_client.delete(f"cache:user:{user_id}")
    except Exception:  # noqa: BLE001 — cache invalidation is best-effort
        pass


@router.get("", response_model=dict, dependencies=[Depends(require_role("admin"))])
@router.get("/", response_model=dict, dependencies=[Depends(require_role("admin"))])
def list_users_route(
    offset: int = Query(0, ge=0),
    limit: int = Query(100, ge=1, le=200),
    db: Session = Depends(get_db),
):
    total = db.query(User).count()
    users = (
        db.query(User)
        .order_by(User.created_at.desc())
        .offset(offset)
        .limit(limit)
        .all()
    )
    items = [_serialize_user(u) for u in users]
    _attach_profiles(items, db)
    return {
        "items": items,
        "total": total,
        "offset": offset,
        "limit": limit,
    }


@router.get("/export", dependencies=[Depends(require_role("admin"))])
def export_users_xlsx(db: Session = Depends(get_db)):
    """Export every user to a .xlsx spreadsheet."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as e:  # pragma: no cover — dependency guarded at deploy time
        raise HTTPException(status_code=503, detail="XLSX export is not available") from e

    users = db.query(User).order_by(User.created_at.desc()).all()
    items = [_serialize_user(u) for u in users]
    _attach_profiles(items, db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    ws.append(EXPORT_HEADERS)

    header_fill = PatternFill(start_color="FF2563EB", end_color="FF2563EB", fill_type="solid")
    header_font = Font(color="FFFFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for item in items:
        profile = item.get("profile") or {}
        ws.append(
            [
                item.get("full_name") or "",
                item.get("email") or "",
                item.get("phone") or "",
                item.get("role") or "",
                "Active" if item.get("is_active") else "Inactive",
                profile.get("form_level") or "",
                profile.get("subjects") or "",
                item.get("created_at") or "",
            ]
        )

    for column_cells in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 4, 10), 40)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.read(),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": 'attachment; filename="casuya-users.xlsx"'},
    )


@router.get("/me", response_model=UserResponse)
@router.get("/me/", response_model=UserResponse)
def get_current_user_route(current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.id == current_user["sub"]).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return UserResponse(
        id=user.id,
        email=user.email,
        full_name=user.full_name,
        phone=user.phone,
        role=user.role,
        is_active=user.is_active,
    )


@router.patch("/me", response_model=UserResponse)
@router.patch("/me/", response_model=UserResponse)
def update_current_user_route(body: UserUpdateRequest, current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.id == current_user["sub"]).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    if body.full_name is not None:
        user.full_name = body.full_name
    if body.phone is not None:
        user.phone = body.phone
    db.commit()
    return UserResponse(
        id=user.id,
        email=user.email,
        full_name=user.full_name,
        phone=user.phone,
        role=user.role,
        is_active=user.is_active,
    )


@router.patch("/{user_id}", response_model=dict, dependencies=[Depends(require_role("admin"))])
def update_user_status(
    user_id: str,
    body: UserStatusUpdateRequest,
    admin=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Activate or deactivate a user account (admin only)."""
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    if not body.is_active and user.id == admin.get("sub"):
        raise HTTPException(status_code=400, detail="You cannot deactivate your own account")
    user.is_active = body.is_active
    db.commit()
    _invalidate_user_cache(user.id)
    return _serialize_user(user)
